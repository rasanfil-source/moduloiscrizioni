from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

base = Path(__file__).resolve().parent
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin = Cm(1.8), Cm(1.7)
s.left_margin = s.right_margin = Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name]
    st.font.name='Calibri'
    st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal']
normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.04
for name,size in [('Title',27),('Heading 1',20),('Heading 2',12)]:
    st=doc.styles[name]
    st.font.size=Pt(size)
    st.paragraph_format.space_before=Pt(10)
    st.paragraph_format.space_after=Pt(6)
doc.styles['Subtitle'].font.size=Pt(13)
pages=base.joinpath('contenuto.txt').read_text(encoding='utf-8').split('===PAGE===')
for idx,page in enumerate(pages):
    if idx: doc.add_page_break()
    for li,line in enumerate(page.strip().splitlines()):
        line=line.strip()
        if not line: continue
        if idx==0 and li==0: doc.add_paragraph(line,'Title')
        elif idx==0 and li==1: doc.add_paragraph(line,'Subtitle')
        elif line.startswith('# '): doc.add_paragraph(line[2:],'Heading 1')
        elif line.startswith('## '): doc.add_paragraph(line[3:],'Heading 2')
        else:
            p=doc.add_paragraph(line)
            if line.startswith('- '):
                p.text='• '+line[2:]
                p.paragraph_format.left_indent=Cm(.3)
                p.paragraph_format.first_line_indent=Cm(-.3)
            p.paragraph_format.widow_control=True
f=s.footer.paragraphs[0]
f.alignment=2
r=f.add_run('Segreteria eventi  •  ')
r.font.size=Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)
doc.core_properties.title='Manuale del segretario per il portale Segreteria eventi'
doc.core_properties.subject='Criteri e procedure operative'
doc.core_properties.author='Segreteria eventi'
for el in list(doc.styles.element.iter(qn('w:pBdr'))):
    el.getparent().remove(el)
for el in list(doc.element.iter(qn('w:pBdr'))):
    el.getparent().remove(el)
doc.save(base/'Manuale_del_segretario.docx')
print(base/'Manuale_del_segretario.docx')
